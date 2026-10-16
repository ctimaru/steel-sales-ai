"""NET-UX3A: pin to ISTAT's official current Italian municipality XLSX.

Only the official URL is fetched; no fallback mirrors or synthesized municipality codes.
Preview validates all rows and a complete 2026 Sardinia/region hierarchy without
using credentials. Apply is an explicitly gated, atomic service-role import.
"""
from __future__ import annotations

import argparse
import asyncio
import hashlib
import io
import json
import re
import unicodedata
import urllib.request
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from app.repository import WorkerRepository
from app.istat_import_legacy import load_verified_csv, normalize_name

SOURCE = "https://www.istat.it/storage/codici-unita-amministrative/Elenco-comuni-italiani.xlsx"
REFERENCE_DATE = date(2026, 2, 21)
EXPECTED_COUNT = 7894
REGIONS = {
    "01": "Piemonte", "02": "Valle d'Aosta", "03": "Lombardia",
    "04": "Trentino-Alto Adige", "05": "Veneto", "06": "Friuli-Venezia Giulia",
    "07": "Liguria", "08": "Emilia-Romagna", "09": "Toscana",
    "10": "Umbria", "11": "Marche", "12": "Lazio", "13": "Abruzzo",
    "14": "Molise", "15": "Campania", "16": "Puglia", "17": "Basilicata",
    "18": "Calabria", "19": "Sicilia", "20": "Sardegna",
}
# Audit exact ISTAT header patterns, avoiding heuristic column guesses.
# Italian 2026 workbook uses official labels and may include bilingual headings.
HEADERS = {
    "istat_code": ("codice comune formato alfanumerico", "codice comune formato numerico"),
    "municipality_name": ("denominazione in italiano", "denominazione comune"),
    "province_code": ("sigla automobilistica",),
    "province_name": ("denominazione dell unita territoriale sovracomunale",
                      "denominazione provincia", "denominazione dell unita territoriale"),
    "region_code": ("codice regione",),
    "region_name": ("denominazione regione",),
}


def _h(value: object) -> str:
    value = unicodedata.normalize("NFKD", str(value or "").lower())
    value = "".join(c for c in value if not unicodedata.combining(c))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", value).split())


def _cells(row: tuple[object, ...], indexes: dict[str, int]) -> dict[str, str]:
    return {field: str(row[idx] if idx < len(row) and row[idx] is not None else "").strip()
            for field, idx in indexes.items()}


def _discover_header(worksheet) -> tuple[int, dict[str, int]]:
    for number, row in enumerate(worksheet.iter_rows(min_row=1,max_row=15,values_only=True),1):
        normalized = [_h(x) for x in row]
        selected: dict[str,int] = {}
        for field, options in HEADERS.items():
            candidates = [i for i,head in enumerate(normalized)
                          if any(head.startswith(option) for option in options)]
            if len(candidates)==1:
                selected[field]=candidates[0]
        if len(selected)==len(HEADERS) and len(set(selected.values()))==len(HEADERS):
            return number, selected
    sample = [[str(x)[:100] for x in row[:8]] for row in
              worksheet.iter_rows(min_row=1,max_row=3,values_only=True)]
    raise ValueError(f"ISTAT 2026 official XLSX schema requires reviewed header mapping; sample={sample!r}")


def load_verified_xlsx(contents: bytes) -> list[dict[str, str]]:
    if len(contents) < 50_000 or not contents.startswith(b"PK"):
        raise ValueError("ISTAT workbook is not a complete XLSX ZIP archive")
    workbook = load_workbook(io.BytesIO(contents),read_only=True,data_only=True)
    try:
        sheet = workbook.active
        header_row, index = _discover_header(sheet)
        records: list[dict[str,str]]=[]
        codes: set[str]=set()
        code_region: dict[str,str]={}
        for row_no, cells in enumerate(sheet.iter_rows(min_row=header_row+1,values_only=True),
                                       header_row+1):
            fields=_cells(cells,index)
            if not any(fields.values()):
                continue
            code=fields["istat_code"].removesuffix(".0").zfill(6)
            region_code=fields["region_code"].removesuffix(".0").zfill(2)
            city=fields["municipality_name"]
            province=fields["province_code"].upper()
            prov_name=fields["province_name"]
            region_name=fields["region_name"]
            if not re.fullmatch(r"[0-9]{6}",code) or code in codes:
                raise ValueError(f"invalid or duplicate ISTAT code at row {row_no}: {code}")
            if region_code not in REGIONS or not re.fullmatch(r"[A-Z]{2}",province):
                raise ValueError(f"invalid region or province code at ISTAT row {row_no}")
            if not city or not prov_name or not region_name:
                raise ValueError(f"missing municipality/province/region at ISTAT row {row_no}")
            if code_region.setdefault(region_code,_h(region_name)) != _h(region_name):
                raise ValueError(f"mixed ISTAT region spellings in code {region_code}")
            # Do not put unofficial bilingual spellings in the Network region filter.
            # ISTAT region CODE is authoritative; canonical Network display is stable.
            codes.add(code)
            records.append({
                "istat_code":code,"municipality_name":city,
                "municipality_normalized":normalize_name(city),
                "province_code":province,"province_name":prov_name,
                "region_code":region_code,"region_name":REGIONS[region_code],
            })
    finally:
        workbook.close()
    if len(records)!=EXPECTED_COUNT or len(code_region)!=20:
        raise ValueError(f"official ISTAT row count or regions mismatch: {len(records)} rows, {len(code_region)} regions")
    caste=[r for r in records if r["municipality_name"].casefold()=="castegnero nanto"]
    if len(caste)!=1 or caste[0]["region_code"]!="05" or caste[0]["province_code"]!="VI":
        raise ValueError("2026 ISTAT municipality Castegnero Nanto (Veneto, VI) missing")
    sardi=[r for r in records if r["region_code"]=="20"]
    if not sardi or not {"112","113","114","115","116","117","118","119"}.issubset(
        {r["istat_code"][:3] for r in sardi}
    ):
        raise ValueError("2026 Sardinia provincial code remapping incomplete")
    return records


def download_official() -> bytes:
    request=urllib.request.Request(SOURCE,headers={"User-Agent":"SmartSteelSales-ISTAT-Reference/1.0"})
    with urllib.request.urlopen(request,timeout=45) as reply:
        if not reply.geturl().startswith("https://www.istat.it/"):
            raise ValueError("ISTAT download unexpectedly redirected off the official host")
        data=reply.read(8_000_001)
    if len(data)>8_000_000:
        raise ValueError("official workbook exceeds 8 MB safety limit")
    return data


async def apply_atomic(records: list[dict[str,str]],fingerprint: str)->dict:
    repo=WorkerRepository()
    header=await repo._request("POST","/rest/v1/network_it_istat_import_batches",json={
        "source_url":SOURCE,"file_sha256":fingerprint,
        "reference_date":REFERENCE_DATE.isoformat(),"expected_count":EXPECTED_COUNT,
    },prefer="return=representation")
    if not isinstance(header,list) or len(header)!=1:
        raise RuntimeError("ISTAT batch create failed: expected one returned batch")
    batch_id=header[0]["id"]
    for start in range(0,len(records),200):
        chunk=[dict(row,batch_id=batch_id) for row in records[start:start+200]]
        await repo._request("POST","/rest/v1/network_it_istat_import_staging",
                            json=chunk,prefer="return=minimal")
    result=await repo._request("POST","/rest/v1/rpc/netux3a_activate_istat_import",
                               json={"p_batch_id":batch_id})
    if not isinstance(result,dict) or result.get("status")!="activated" or result.get("municipalities")!=7894:
        raise RuntimeError(f"ISTAT activation returned unexpected result: {result!r}")
    return result


def main()->None:
    parser=argparse.ArgumentParser(description="Official ISTAT 2026 XLSX guarded load")
    parser.add_argument("--mode",choices=("preview","apply"),default="preview")
    parser.add_argument("--file",type=Path,help="Optional verified XLSX from the official permalink")
    parser.add_argument("--expected-sha256",help="Optional pinned workbook SHA-256")
    args=parser.parse_args()
    if args.mode=="apply" and args.file:
        parser.error("apply mode must re-download directly from ISTAT; no user-supplied file")
    data=args.file.read_bytes() if args.file else download_official()
    sha=hashlib.sha256(data).hexdigest()
    if args.expected_sha256 and sha!=args.expected_sha256.lower():
        raise ValueError("ISTAT official XLSX SHA-256 changed from approved checksum")
    records=load_verified_xlsx(data)
    print(json.dumps({"status":"validated","count":len(records),
                      "reference_date":REFERENCE_DATE.isoformat(),"sha256":sha,
                      "source":SOURCE,"regions":len(REGIONS),"apply":args.mode=="apply"},
                     ensure_ascii=False),flush=True)
    if args.mode=="apply":
        print(json.dumps(asyncio.run(apply_atomic(records,sha))),flush=True)


if __name__=="__main__":
    main()
