from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass
from email import policy
from email.message import Message
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
from typing import Any

MAX_CHUNK_CHARS = 1800
CHUNK_OVERLAP_CHARS = 180

ITALIAN_MARKERS = {
    "il", "la", "di", "che", "per", "con", "non", "una", "disponibile",
    "disponibilita", "consegna", "produzione", "prezzo", "offerta",
}
ENGLISH_MARKERS = {
    "the", "and", "of", "for", "with", "not", "available", "availability",
    "delivery", "production", "price", "offer", "quotation",
}


@dataclass(frozen=True)
class TextUnit:
    text: str
    locator: dict[str, Any]
    section_path: list[str]
    page: int | None = None


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def detect_language(text: str) -> str:
    words = re.findall(r"[a-zA-ZÀ-ÿ]+", text.casefold())
    if not words:
        return "und"
    it_score = sum(word in ITALIAN_MARKERS for word in words)
    en_score = sum(word in ENGLISH_MARKERS for word in words)
    if it_score == en_score == 0:
        return "und"
    if it_score >= en_score + 2:
        return "it"
    if en_score >= it_score + 2:
        return "en"
    return "und"


def _email_text(message: Message) -> str:
    parts: list[str] = []
    subject = str(message.get("subject") or "").strip()
    if subject:
        parts.append(f"Subject: {subject}")
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_type() != "text/plain":
                continue
            disposition = str(part.get("Content-Disposition") or "").casefold()
            if "attachment" in disposition:
                continue
            try:
                parts.append(str(part.get_content()))
            except Exception:
                continue
    else:
        try:
            parts.append(str(message.get_content()))
        except Exception:
            pass
    return "\n\n".join(part.strip() for part in parts if part and part.strip())


def _email_unit(payload: bytes, *, archive_member: str | None = None) -> TextUnit:
    message = BytesParser(policy=policy.default).parsebytes(payload)
    subject = str(message.get("subject") or "").strip()
    locator: dict[str, Any] = {
        "kind": "email",
        "message_id": str(message.get("message-id") or "").strip() or None,
        "subject": subject or None,
        "from": str(message.get("from") or "").strip() or None,
        "to": str(message.get("to") or "").strip() or None,
        "date": str(message.get("date") or "").strip() or None,
    }
    if archive_member:
        locator["archive_member"] = archive_member
    path = ["Email"]
    if archive_member:
        path = ["Archive", archive_member, "Email"]
    if subject:
        path.append(subject[:160])
    return TextUnit(text=_email_text(message), locator=locator, section_path=path)


def extract_text_units(filename: str, payload: bytes) -> list[TextUnit]:
    extension = Path(filename).suffix.casefold()
    if extension == ".eml":
        return [_email_unit(payload)]

    if extension == ".zip":
        units: list[TextUnit] = []
        try:
            with zipfile.ZipFile(BytesIO(payload)) as archive:
                for member in sorted(archive.infolist(), key=lambda item: item.filename.casefold()):
                    if member.is_dir():
                        continue
                    content = archive.read(member)
                    name = member.filename
                    lower = name.casefold()
                    if lower.endswith(".eml"):
                        units.append(_email_unit(content, archive_member=name))
                    elif lower.endswith((".txt", ".csv", ".md")):
                        text = content.decode("utf-8", errors="replace").strip()
                        if text:
                            units.append(
                                TextUnit(
                                    text=text,
                                    locator={"kind": "archive_text", "archive_member": name},
                                    section_path=["Archive", name],
                                )
                            )
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid ZIP archive.") from exc
        return units

    if extension == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(payload))
            return [
                TextUnit(
                    text=(page.extract_text() or "").strip(),
                    locator={"kind": "pdf_page", "page": index},
                    section_path=["PDF", f"Page {index}"],
                    page=index,
                )
                for index, page in enumerate(reader.pages, start=1)
                if (page.extract_text() or "").strip()
            ]
        except Exception as exc:
            raise ValueError("Invalid PDF document.") from exc

    if extension == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
            units: list[TextUnit] = []
            for sheet in workbook.worksheets:
                lines = []
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                    if values:
                        lines.append(f"R{row_number}: " + " | ".join(values))
                if lines:
                    units.append(
                        TextUnit(
                            text="\n".join(lines),
                            locator={"kind": "spreadsheet_sheet", "sheet": sheet.title},
                            section_path=["Workbook", sheet.title],
                        )
                    )
            return units
        except Exception as exc:
            raise ValueError("Invalid XLSX document.") from exc

    if extension == ".xls":
        try:
            import xlrd

            workbook = xlrd.open_workbook(file_contents=payload)
            units = []
            for sheet in workbook.sheets():
                lines = []
                for row_number in range(sheet.nrows):
                    values = [str(value).strip() for value in sheet.row_values(row_number) if str(value).strip()]
                    if values:
                        lines.append(f"R{row_number + 1}: " + " | ".join(values))
                if lines:
                    units.append(
                        TextUnit(
                            text="\n".join(lines),
                            locator={"kind": "spreadsheet_sheet", "sheet": sheet.name},
                            section_path=["Workbook", sheet.name],
                        )
                    )
            return units
        except Exception as exc:
            raise ValueError("Invalid XLS document.") from exc

    text = payload.decode("utf-8", errors="replace").strip()
    return [TextUnit(text=text, locator={"kind": "text"}, section_path=["Document"])] if text else []


def _split_unit(unit: TextUnit) -> list[tuple[str, int, int]]:
    text = unit.text
    if not text:
        return []
    chunks: list[tuple[str, int, int]] = []
    start = 0
    while start < len(text):
        end = min(start + MAX_CHUNK_CHARS, len(text))
        if end < len(text):
            break_at = text.rfind("\n", start + MAX_CHUNK_CHARS // 2, end)
            if break_at < 0:
                break_at = text.rfind(" ", start + MAX_CHUNK_CHARS // 2, end)
            if break_at > start:
                end = break_at
        left = start
        while left < end and text[left].isspace():
            left += 1
        right = end
        while right > left and text[right - 1].isspace():
            right -= 1
        if right > left:
            chunks.append((text[left:right], left, right))
        if end >= len(text):
            break
        start = max(end - CHUNK_OVERLAP_CHARS, start + 1)
    return chunks


def build_knowledge_document(
    *,
    filename: str,
    payload: bytes,
    storage_path: str,
    document_type: str,
    extraction_version: str,
) -> dict[str, Any]:
    units = extract_text_units(filename, payload)
    chunks: list[dict[str, Any]] = []
    languages: list[str] = []
    for unit_index, unit in enumerate(units):
        language = detect_language(unit.text)
        if language != "und":
            languages.append(language)
        for content, start, end in _split_unit(unit):
            chunk_language = detect_language(content)
            locator = dict(unit.locator)
            locator.update(
                {
                    "unit_index": unit_index,
                    "unit_char_start": start,
                    "unit_char_end": end,
                }
            )
            chunk: dict[str, Any] = {
                "chunk_index": len(chunks),
                "content": content,
                "content_checksum": sha256_text(content),
                "language_code": chunk_language,
                "token_count": len(content.split()),
                "section_path": unit.section_path,
                "source_locator": locator,
                "metadata": {"chunking": "knowledge-v1", "overlap_chars": CHUNK_OVERLAP_CHARS},
            }
            if unit.page is not None:
                chunk["page_start"] = unit.page
                chunk["page_end"] = unit.page
            chunks.append(chunk)

    language_code = max(set(languages), key=languages.count) if languages else "und"
    title = filename
    if units and units[0].locator.get("subject"):
        title = str(units[0].locator["subject"])
    mime_type = {
        ".eml": "message/rfc822",
        ".zip": "application/zip",
        ".pdf": "application/pdf",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xls": "application/vnd.ms-excel",
    }.get(Path(filename).suffix.casefold(), "application/octet-stream")

    return {
        "external_id": filename,
        "document_type": document_type,
        "title": title,
        "filename": filename,
        "mime_type": mime_type,
        "storage_path": storage_path,
        "content_checksum": sha256_bytes(payload),
        "extraction_method": "steel-sales-ai-worker",
        "extraction_version": extraction_version,
        "language_code": language_code,
        "chunks": chunks,
    }
