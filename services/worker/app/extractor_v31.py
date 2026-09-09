from __future__ import annotations

import re
import zipfile
from email import policy
from email.parser import BytesParser
from io import BytesIO
from typing import Any

NUMBER = r"\d+(?:[,.]\d+)?"
DIMENSIONS_RE = re.compile(
    rf"(?P<a>{NUMBER})\s*[x×]\s*(?P<b>{NUMBER})(?:\s*[x×]\s*(?P<c>{NUMBER}))?",
    re.IGNORECASE,
)
GRADE_RE = re.compile(r"\b(?:P265GH|P355(?:GH|NH)|S\d{3}(?:J2H?|JR)?|L275)\b", re.IGNORECASE)
STANDARD_RE = re.compile(r"\bEN\s*\d{4,5}(?:\s*[A-Z]\d+)?\b", re.IGNORECASE)
PRICE_RE = re.compile(
    rf"(?:€|EUR\s*)\s*(?P<value>{NUMBER})\s*(?:/\s*|per\s*)?(?P<unit>mt|m|kg)?",
    re.IGNORECASE,
)
PRICE_SUFFIX_RE = re.compile(
    rf"(?P<value>{NUMBER})\s*(?:€|EUR)\s*(?:/\s*)?(?P<unit>mt|m|kg)?",
    re.IGNORECASE,
)
QUANTITY_RE = re.compile(rf"(?P<value>{NUMBER})\s*(?P<unit>ton(?:nellate)?|t|kg)\b", re.IGNORECASE)
LENGTH_RE = re.compile(rf"(?P<value>{NUMBER})\s*(?P<unit>mm|mt|m)\b", re.IGNORECASE)

NEGATIVE_AVAILABILITY = (
    "non disponibile",
    "non ce la facciamo",
    "non abbiamo",
    "nessuna disponibil",
    "non è disponibile",
)
POSITIVE_AVAILABILITY = ("disponibile", "a terra", "stock", "pronta consegna")
PRODUCTION_AVAILABILITY = ("produzione", "laminazione", "rientrare", "rientro")


def number(value: str) -> float:
    normalized = value.replace(".", "").replace(",", ".")
    return float(normalized)


def scalar(value: float) -> int | float:
    return int(value) if value.is_integer() else value


def _message_text(payload: bytes) -> str:
    message = BytesParser(policy=policy.default).parsebytes(payload)
    parts = [f"[SUBJECT] {message.get('subject', '')}"]
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                parts.append(part.get_content())
    else:
        parts.append(message.get_content())
    return "\n".join(str(part) for part in parts if part)


def _document_text(filename: str, payload: bytes) -> str:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if extension == "eml":
        return _message_text(payload)
    if extension == "zip":
        parts: list[str] = []
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            for member in archive.infolist():
                if member.is_dir():
                    continue
                name = member.filename.lower()
                content = archive.read(member)
                if name.endswith(".eml"):
                    parts.append(_message_text(content))
                elif name.endswith((".txt", ".csv", ".md")):
                    parts.append(content.decode("utf-8", errors="replace"))
        return "\n".join(parts)
    if extension == "pdf":
        try:
            from pypdf import PdfReader

            return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(payload)).pages)
        except ImportError as exc:
            raise RuntimeError("PDF parsing requires the pypdf dependency.") from exc
    if extension == "xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
            return "\n".join(
                " | ".join(str(cell) for cell in row if cell is not None)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(values_only=True)
                if any(cell is not None for cell in row)
            )
        except ImportError as exc:
            raise RuntimeError("XLSX parsing requires the openpyxl dependency.") from exc
    if extension == "xls":
        try:
            import xlrd

            workbook = xlrd.open_workbook(file_contents=payload)
            return "\n".join(
                " | ".join(str(value) for value in sheet.row_values(row))
                for sheet in workbook.sheets()
                for row in range(sheet.nrows)
            )
        except ImportError as exc:
            raise RuntimeError("XLS parsing requires the xlrd dependency.") from exc
    return payload.decode("utf-8", errors="replace")


def _role(line: str, has_price: bool) -> str:
    lowered = line.casefold()
    if any(word in lowered for word in ("consegnato", "consegnata", "consegna effettuata")):
        return "delivered"
    if any(word in lowered for word in ("ordine", "ordinato", "ordinata", "confermato ordine")):
        return "ordered"
    if has_price or any(word in lowered for word in ("offro", "offerta", "disponibile", "a terra")):
        return "offered"
    if any(word in lowered for word in ("richiesta", "servono", "cerco", "quanto te ne serve", "novità sulla produzione")):
        return "requested"
    return "requested"


def _availability(line: str) -> str | None:
    lowered = line.casefold()
    if any(term in lowered for term in NEGATIVE_AVAILABILITY):
        return "unavailable"
    if any(term in lowered for term in POSITIVE_AVAILABILITY):
        return "stock"
    if any(term in lowered for term in PRODUCTION_AVAILABILITY):
        return "production"
    return None


def extract_observations(text: str, source_filename: str) -> list[dict[str, Any]]:
    observations: list[dict[str, Any]] = []
    for raw_line in text.splitlines():
        line = " ".join(raw_line.split()).strip()
        if not line:
            continue
        dimensions = DIMENSIONS_RE.search(line)
        grade_match = GRADE_RE.search(line)
        standard_match = STANDARD_RE.search(line)
        price_match = PRICE_RE.search(line) or PRICE_SUFFIX_RE.search(line)
        quantity_match = QUANTITY_RE.search(line)
        length_match = LENGTH_RE.search(line)
        if not any((dimensions, grade_match, standard_match, price_match, quantity_match)):
            continue

        row: dict[str, Any] = {
            "source_filename": source_filename,
            "source_text": line,
            "item_role": _role(line, price_match is not None),
            "grade": grade_match.group(0).upper() if grade_match else None,
            "standard": standard_match.group(0).upper().replace(" ", " ") if standard_match else None,
            "availability_status": _availability(line),
            "confidence": 0.98 if dimensions and grade_match else 0.88 if dimensions else 0.72,
        }
        if dimensions:
            first, second = number(dimensions.group("a")), number(dimensions.group("b"))
            third = number(dimensions.group("c")) if dimensions.group("c") else None
            if abs(first - second) < 0.001 and third is not None:
                row.update(width_mm=scalar(first), height_mm=scalar(second), thickness_mm=scalar(third))
            else:
                row.update(outer_diameter_mm=scalar(first), thickness_mm=scalar(second))
                if third is not None:
                    row["length_mm"] = scalar(third)
        if length_match and "length_mm" not in row:
            length = number(length_match.group("value"))
            row["length_mm"] = scalar(length * 1000 if length_match.group("unit").lower() in ("m", "mt") and length < 100 else length)
        if quantity_match:
            row["quantity"] = scalar(number(quantity_match.group("value")))
            unit = quantity_match.group("unit").lower()
            row["quantity_unit"] = "T" if unit in ("t", "ton", "tonnellate") else unit.upper()
        if price_match:
            row["price_value"] = scalar(number(price_match.group("value")))
            row["price_unit"] = (price_match.group("unit") or "M").upper()
            row["currency"] = "EUR"
        observations.append(row)
    return observations


def parse_document(filename: str, payload: bytes) -> list[dict[str, Any]]:
    return extract_observations(_document_text(filename, payload), filename)
