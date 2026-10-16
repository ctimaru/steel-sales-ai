from __future__ import annotations

import io

from openpyxl import Workbook
import pytest

from app.istat_import import (
    EXPECTED_COUNT, REGIONS, SOURCE, _discover_header, load_verified_xlsx,
)


def test_istat_header_mapping_is_explicit() -> None:
    wb = Workbook()
    sh = wb.active
    sh.append([
        "Codice Comune formato alfanumerico",
        "Denominazione in italiano",
        "Sigla automobilistica",
        "Denominazione dell’Unità territoriale sovracomunale (valida a fini statistici)",
        "Codice Regione",
        "Denominazione Regione",
    ])
    nr, ix = _discover_header(sh)
    assert nr == 1
    assert set(ix) == {
        "istat_code", "municipality_name", "province_code",
        "province_name", "region_code", "region_name",
    }


def test_istat_schema_fails_closed_on_unknown_headers() -> None:
    wb = Workbook()
    wb.active.append(["COMUNE", "NOME", "REGIONE"])
    with pytest.raises(ValueError, match="reviewed header mapping"):
        _discover_header(wb.active)


def test_full_synthetic_workbook_contract_and_2026_sardinia() -> None:
    # Purely in-memory fixture: synthetic ISTAT-shaped rows are NEVER staged or
    # written to Supabase. This is a validation contract test, not reference data.
    wb = Workbook(write_only=True)
    sh = wb.create_sheet("Comuni")
    sh.append([
        "Codice Comune formato alfanumerico",
        "Denominazione in italiano",
        "Sigla automobilistica",
        "Denominazione dell'Unità territoriale sovracomunale",
        "Codice Regione", "Denominazione Regione",
    ])
    sardinia_code = 112
    for i in range(EXPECTED_COUNT):
        region_code = f"{i % 20 + 1:02d}"
        code = f"{i + 1:06d}"
        city = f"Comune sintetico #{i}"
        province = "MI"
        if i == 4:
            city, province = "Castegnero Nanto", "VI"
        if region_code == "20" and sardinia_code <= 119:
            code = f"{sardinia_code:03d}001"
            sardinia_code += 1
        sh.append([code, city, province, "Provincia test", region_code, REGIONS[region_code]])
    sink = io.BytesIO()
    wb.save(sink)
    contents = sink.getvalue()
    assert len(contents) > 50_000
    result = load_verified_xlsx(contents)
    assert len(result) == 7894
    assert len(set(r["istat_code"] for r in result)) == 7894
    assert len(set(r["region_code"] for r in result)) == 20
    assert SOURCE.startswith("https://www.istat.it/")
